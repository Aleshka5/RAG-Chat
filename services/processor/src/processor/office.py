from __future__ import annotations

import zipfile
from dataclasses import dataclass, field
from io import BytesIO

from docx import Document
from openpyxl import load_workbook

_HEADING_NAMES = frozenset({"heading 1", "heading 2", "heading 3"})
_HEADING_IDS = frozenset({"heading1", "heading2", "heading3"})


class CorruptDocument(Exception):
    pass


@dataclass(frozen=True)
class OfficeUnit:
    label: str
    text: str


@dataclass
class _Draft:
    label: str | None
    lines: list[str] = field(default_factory=list)


def extract_docx(data: bytes) -> list[OfficeUnit]:
    names = _zip_names(data)
    if "word/document.xml" not in names:
        raise CorruptDocument("missing document.xml")
    try:
        document = Document(BytesIO(data))
        return _units_from_docx(document)
    except (KeyboardInterrupt, SystemExit):
        raise
    except Exception as exc:
        raise CorruptDocument("unreadable") from exc


def extract_xlsx(data: bytes) -> list[OfficeUnit]:
    names = _zip_names(data)
    if "xl/workbook.xml" not in names:
        raise CorruptDocument("missing workbook.xml")
    try:
        workbook = load_workbook(BytesIO(data), data_only=True)
    except (KeyboardInterrupt, SystemExit):
        raise
    except Exception as exc:
        raise CorruptDocument("unreadable") from exc
    try:
        return _units_from_xlsx(workbook)
    except (KeyboardInterrupt, SystemExit):
        raise
    except Exception as exc:
        raise CorruptDocument("unreadable") from exc
    finally:
        workbook.close()


def _zip_names(data: bytes) -> set[str]:
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            return set(archive.namelist())
    except (KeyboardInterrupt, SystemExit):
        raise
    except Exception as exc:
        raise CorruptDocument("unreadable") from exc


def _is_heading(paragraph: object) -> bool:
    style = getattr(paragraph, "style", None)
    if style is None:
        return False
    name = style.name.strip().lower() if isinstance(style.name, str) else ""
    style_id = style.style_id.strip().lower() if isinstance(getattr(style, "style_id", None), str) else ""
    return name in _HEADING_NAMES or style_id in _HEADING_IDS


def _units_from_docx(document: object) -> list[OfficeUnit]:
    drafts: list[_Draft] = []

    def current() -> _Draft:
        if not drafts:
            drafts.append(_Draft(label=None))
        return drafts[-1]

    for paragraph in document.paragraphs:
        text = paragraph.text or ""
        if _is_heading(paragraph):
            drafts.append(_Draft(label=text.strip()))
            if text.strip():
                drafts[-1].lines.append(text)
            continue
        if text.strip() == "":
            continue
        current().lines.append(text)

    if document.tables:
        target = current()
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text for cell in row.cells]
                if all(cell.strip() == "" for cell in cells):
                    continue
                target.lines.append(" | ".join(cells))

    return _finish(drafts)


def _units_from_xlsx(workbook: object) -> list[OfficeUnit]:
    units: list[OfficeUnit] = []
    for sheet in workbook.worksheets:
        lines: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            rendered = ["" if value is None else str(value) for value in row]
            if all(value.strip() == "" for value in rendered):
                continue
            lines.append("\t".join(rendered))
        if not lines:
            continue
        units.append(OfficeUnit(label=sheet.title, text="\n".join(lines)))
    return units


def _finish(drafts: list[_Draft]) -> list[OfficeUnit]:
    kept = [draft for draft in drafts if "\n".join(draft.lines).strip()]
    units: list[OfficeUnit] = []
    for index, draft in enumerate(kept, start=1):
        label = draft.label if draft.label else f"Раздел {index}"
        units.append(OfficeUnit(label=label, text="\n".join(draft.lines)))
    return units
