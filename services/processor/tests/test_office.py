from __future__ import annotations

import zipfile
from io import BytesIO

import pytest
from docx import Document
from openpyxl import Workbook
from tests.fakes import FakeQdrant, MemoryBlobs, MemoryFiles
from tests.support import build_pipeline, file_document, make_settings, message_for

from processor.messages import DOCX_CONTENT_TYPE, XLSX_CONTENT_TYPE
from processor.office import OfficeUnit, extract_docx, extract_xlsx

SECTION = "абзац один " * 8
LEFT = "лево " * 8
RIGHT = "значение " * 8
PAGE_FIELDS = {"file_id", "user_id", "page_number", "text", "origin", "updated_at"}
FIRST_UNIT = f"Первый\n{SECTION}\nПодпункт"
SECOND_UNIT = "Второй\nабзац два\na | b\nc | d"
FIRST_SHEET = f"{LEFT}\t2\nниз\tкрай"
SECOND_SHEET = f"{RIGHT}\t\tхвост"


def _docx_bytes(build) -> bytes:
    document = Document()
    build(document)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes(build) -> bytes:
    workbook = Workbook()
    build(workbook)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _two_section_docx() -> bytes:
    def build(document: Document) -> None:
        document.add_heading("Первый", level=1)
        document.add_paragraph(SECTION)
        document.add_paragraph("")
        document.add_paragraph("   ")
        document.add_heading("Подпункт", level=4)
        document.add_heading("Второй", level=1)
        document.add_paragraph("абзац два")
        table = document.add_table(rows=3, cols=2)
        table.cell(0, 0).text = "a"
        table.cell(0, 1).text = "b"
        table.cell(1, 0).text = " "
        table.cell(1, 1).text = ""
        table.cell(2, 0).text = "c"
        table.cell(2, 1).text = "d"

    return _docx_bytes(build)


def _two_sheet_xlsx() -> bytes:
    def build(workbook: Workbook) -> None:
        first = workbook.active
        first.title = "Первый"
        first["A1"] = LEFT
        first["B1"] = 2
        first["A3"] = "низ"
        first["B3"] = "край"
        workbook.create_sheet("Пусто")
        second = workbook.create_sheet("Второй")
        second["A1"] = RIGHT
        second["B1"] = "=SUM(A1)"
        second["C1"] = "хвост"

    return _xlsx_bytes(build)


def _plain_docx() -> bytes:
    return _docx_bytes(lambda document: document.add_paragraph("просто текст"))


def _empty_docx() -> bytes:
    return _docx_bytes(lambda _document: None)


def _empty_xlsx() -> bytes:
    return _xlsx_bytes(lambda _workbook: None)


def _foreign_zip() -> bytes:
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("hello.txt", "hi")
    return buffer.getvalue()


def _replace_in_zip(payload: bytes, old: bytes, new: bytes) -> bytes:
    source = zipfile.ZipFile(BytesIO(payload))
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if old in data:
                data = data.replace(old, new, 1)
            target.writestr(info, data)
    source.close()
    return buffer.getvalue()


def test_docx_heading_rules_number_kept_units_and_append_tables_to_the_last() -> None:
    def build(document: Document) -> None:
        renamed = document.styles["Heading 1"]
        renamed.name = "Заголовок 1"
        document.add_paragraph("преамбула")
        document.add_paragraph("Глава", style=renamed)
        document.add_paragraph("тело")
        document.add_heading("   ", level=2)
        document.add_heading("   ", level=3)
        document.add_paragraph("после пустого")
        document.add_heading("Мелкий", level=4)
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "л"
        table.cell(0, 1).text = "п"
        table.cell(1, 0).text = ""
        table.cell(1, 1).text = " "

    units = extract_docx(_docx_bytes(build))
    assert [(unit.label, unit.text) for unit in units] == [
        ("Раздел 1", "преамбула"),
        ("Глава", "Глава\nтело"),
        ("Раздел 3", "после пустого\nМелкий\nл | п"),
    ]


def test_docx_tables_without_paragraph_units_use_section_1() -> None:
    def build(document: Document) -> None:
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "ячейка"
        table.cell(0, 1).text = "два"

    assert extract_docx(_docx_bytes(build)) == [OfficeUnit("Раздел 1", "ячейка | два")]


def test_xlsx_reads_a_cached_formula_value() -> None:
    def build(workbook: Workbook) -> None:
        sheet = workbook.active
        sheet.title = "Calc"
        sheet["A1"] = 1
        sheet["B1"] = "=A1+1"

    payload = _replace_in_zip(_xlsx_bytes(build), b"<f>A1+1</f><v></v>", b"<f>A1+1</f><v>11</v>")
    units = extract_xlsx(payload)
    assert units == [OfficeUnit("Calc", "1\t11")]
    assert "A1+1" not in units[0].text


def test_docx_pipeline_splits_heading_sections_and_skips_ocr() -> None:
    settings = make_settings(chunk_tokens=16, chunk_overlap_tokens=0)
    document = file_document(content_type=DOCX_CONTENT_TYPE, file_name="Отчёт.docx")
    store = MemoryFiles()
    store.add(document)
    blobs = MemoryBlobs(_two_section_docx())
    pipeline, ocr, embedding, qdrant = build_pipeline(settings, store, blobs)
    result = pipeline.handle(message_for(document, content_type=DOCX_CONTENT_TYPE), redelivered=False)
    assert result.action == "ack"
    saved = store.get_by_file_id(document["file_id"])
    assert saved is not None
    assert saved["status"] == "ready"
    assert saved["page_count"] == 2
    assert saved["checkpoint"]["pages_total"] == 2
    assert saved["checkpoint"]["pages_done"] == 2
    pages = store.list_pages(document["file_id"])
    assert [page["text"] for page in pages] == [FIRST_UNIT, SECOND_UNIT]
    assert [page["origin"] for page in pages] == ["text", "text"]
    assert [page["page_number"] for page in pages] == [1, 2]
    assert all(set(page) == PAGE_FIELDS for page in pages)
    assert ocr.calls == []
    assert embedding.image_calls == []
    assert qdrant.points
    labels = {point["payload"]["unit_label"] for point in qdrant.points.values()}
    assert labels == {"Первый", "Второй"}
    for point in qdrant.points.values():
        payload = point["payload"]
        assert payload["source"] == "text"
        assert payload["document_id"] == document["file_id"]
        assert payload["file_name"] == "Отчёт.docx"
        assert isinstance(payload["page_number"], int)


def test_docx_without_headings_is_one_section() -> None:
    settings = make_settings()
    document = file_document(content_type=DOCX_CONTENT_TYPE, file_name="Заметка.docx")
    store = MemoryFiles()
    store.add(document)
    pipeline, ocr, _embedding, qdrant = build_pipeline(settings, store, MemoryBlobs(_plain_docx()))
    assert pipeline.handle(message_for(document, content_type=DOCX_CONTENT_TYPE), redelivered=False).action == "ack"
    saved = store.get_by_file_id(document["file_id"])
    assert saved["page_count"] == 1
    assert saved["checkpoint"]["pages_total"] == 1
    pages = store.list_pages(document["file_id"])
    assert pages[0]["text"] == "просто текст"
    assert pages[0]["origin"] == "text"
    assert ocr.calls == []
    assert {point["payload"]["unit_label"] for point in qdrant.points.values()} == {"Раздел 1"}
    assert all(point["payload"]["source"] == "text" for point in qdrant.points.values())


def test_xlsx_pipeline_keeps_non_empty_sheets_and_drops_uncached_formulas() -> None:
    settings = make_settings(chunk_tokens=16, chunk_overlap_tokens=0)
    document = file_document(content_type=XLSX_CONTENT_TYPE, file_name="Таблица.xlsx")
    store = MemoryFiles()
    store.add(document)
    pipeline, ocr, _embedding, qdrant = build_pipeline(settings, store, MemoryBlobs(_two_sheet_xlsx()))
    result = pipeline.handle(message_for(document, content_type=XLSX_CONTENT_TYPE), redelivered=False)
    assert result.action == "ack"
    saved = store.get_by_file_id(document["file_id"])
    assert saved["status"] == "ready"
    assert saved["page_count"] == 2
    assert saved["checkpoint"]["pages_total"] == 2
    assert saved["checkpoint"]["pages_done"] == 2
    pages = store.list_pages(document["file_id"])
    assert [page["text"] for page in pages] == [FIRST_SHEET, SECOND_SHEET]
    assert [page["origin"] for page in pages] == ["text", "text"]
    assert "=SUM(A1)" not in SECOND_SHEET
    assert all("=SUM" not in page["text"] for page in pages)
    assert ocr.calls == []
    labels = [point["payload"]["unit_label"] for point in qdrant.points.values()]
    assert set(labels) == {"Первый", "Второй"}
    assert "Пусто" not in labels
    assert all(point["payload"]["source"] == "text" for point in qdrant.points.values())


@pytest.mark.parametrize("content_type", [DOCX_CONTENT_TYPE, XLSX_CONTENT_TYPE])
@pytest.mark.parametrize("payload", [b"this is not a pdf", _foreign_zip()], ids=["bytes", "zip"])
def test_unreadable_office_package_is_corrupt_document(content_type: str, payload: bytes) -> None:
    settings = make_settings()
    document = file_document(content_type=content_type, file_name="Сломанный.bin")
    store = MemoryFiles()
    store.add(document)
    pipeline, ocr, _embedding, qdrant = build_pipeline(settings, store, MemoryBlobs(payload))
    result = pipeline.handle(message_for(document, content_type=content_type), redelivered=False)
    assert result.action == "nack"
    assert result.requeue is False
    saved = store.get_by_file_id(document["file_id"])
    assert saved is not None
    assert saved["status"] == "error"
    assert saved["error"] == "corrupt_document"
    assert saved["page_count"] is None
    assert qdrant.points == {}
    assert ("upsert",) not in [(op[0],) for op in qdrant.ops]
    assert ocr.calls == []


@pytest.mark.parametrize(
    ("content_type", "payload"),
    [(DOCX_CONTENT_TYPE, _empty_docx()), (XLSX_CONTENT_TYPE, _empty_xlsx())],
)
def test_office_file_without_text_is_no_extractable_text(content_type: str, payload: bytes) -> None:
    settings = make_settings()
    document = file_document(content_type=content_type, file_name="Пустой.bin")
    store = MemoryFiles()
    store.add(document)
    qdrant = FakeQdrant()
    pipeline, ocr, _embedding, qdrant = build_pipeline(settings, store, MemoryBlobs(payload), qdrant=qdrant)
    result = pipeline.handle(message_for(document, content_type=content_type), redelivered=False)
    assert result.action == "nack"
    assert result.requeue is False
    saved = store.get_by_file_id(document["file_id"])
    assert saved["status"] == "error"
    assert saved["error"] == "no_extractable_text"
    assert ("delete_file", document["user_id"], document["file_id"]) in qdrant.ops
    assert ocr.calls == []


def test_office_resume_does_not_rewrite_a_stored_unit_or_move_pages_done_back() -> None:
    settings = make_settings(chunk_tokens=16, chunk_overlap_tokens=0)
    document = file_document(
        status="processing",
        content_type=DOCX_CONTENT_TYPE,
        file_name="Отчёт.docx",
        checkpoint={"pages_done": 5, "pages_total": 2, "process_attempts": 1},
    )
    store = MemoryFiles()
    store.add(document)
    stored = "STORED UNIT " * 10
    store.insert_page(
        {
            "file_id": document["file_id"],
            "user_id": document["user_id"],
            "page_number": 1,
            "text": stored,
            "origin": "text",
            "updated_at": document["updated_at"],
        }
    )
    inserted: list[int] = []
    original = store.insert_page

    def spy(page: dict) -> None:
        inserted.append(page["page_number"])
        original(page)

    store.insert_page = spy
    pipeline, ocr, _embedding, qdrant = build_pipeline(settings, store, MemoryBlobs(_two_section_docx()))
    result = pipeline.handle(message_for(document, content_type=DOCX_CONTENT_TYPE), redelivered=True)
    assert result.action == "ack"
    assert inserted == [2]
    pages = store.list_pages(document["file_id"])
    assert pages[0]["text"] == stored
    assert pages[1]["text"] == SECOND_UNIT
    assert pages[1]["origin"] == "text"
    assert store.pages_done_history
    assert all(value >= 5 for value in store.pages_done_history)
    assert store.pages_done_history == sorted(store.pages_done_history)
    saved = store.get_by_file_id(document["file_id"])
    assert saved["checkpoint"]["pages_done"] >= 5
    assert saved["page_count"] == 2
    assert ocr.calls == []
    assert any(
        "STORED UNIT" in point["payload"]["text"] and point["payload"]["unit_label"] == "Первый"
        for point in qdrant.points.values()
    )
    assert any(point["payload"]["unit_label"] == "Второй" for point in qdrant.points.values())


def test_office_chunk_setting_change_keeps_file_pages_and_rewrites_points() -> None:
    settings = make_settings(chunk_tokens=8, chunk_overlap_tokens=0)
    document = file_document(
        content_type=DOCX_CONTENT_TYPE,
        file_name="Отчёт.docx",
        checkpoint={
            "embedding_model": settings.embedding_model,
            "chunk_tokens": 32,
            "chunk_overlap": 0,
            "chunks_done": 4,
            "pages_done": 1,
            "pages_total": 1,
        },
    )
    store = MemoryFiles()
    store.add(document)
    stored = "PRESEEDED OFFICE PAGE"
    store.insert_page(
        {
            "file_id": document["file_id"],
            "user_id": document["user_id"],
            "page_number": 1,
            "text": stored,
            "origin": "text",
            "updated_at": document["updated_at"],
        }
    )
    qdrant = FakeQdrant()
    qdrant.points["old"] = {
        "id": "old",
        "vector": [1.0] * settings.embedding_dim,
        "payload": {
            "user_id": document["user_id"],
            "document_id": document["file_id"],
            "chunk_index": 0,
        },
    }
    pipeline, ocr, _embedding, qdrant = build_pipeline(
        settings,
        store,
        MemoryBlobs(_two_section_docx()),
        qdrant=qdrant,
    )
    assert pipeline.handle(message_for(document, content_type=DOCX_CONTENT_TYPE), redelivered=False).action == "ack"
    pages = store.list_pages(document["file_id"])
    assert pages[0]["text"] == stored
    assert pages[0]["origin"] == "text"
    assert "unit_label" not in pages[0]
    assert pages[1]["text"] == SECOND_UNIT
    assert qdrant.ops[0][0] == "delete_file"
    assert any(op[0] == "upsert" for op in qdrant.ops)
    assert "old" not in qdrant.points
    assert qdrant.points
    assert any(
        "OFFICE" in point["payload"]["text"] and point["payload"]["unit_label"] == "Первый"
        for point in qdrant.points.values()
    )
    saved = store.get_by_file_id(document["file_id"])
    assert saved["status"] == "ready"
    assert saved["checkpoint"]["chunk_tokens"] == settings.chunk_tokens
    assert saved["checkpoint"]["chunks_done"] == saved["checkpoint"]["chunks_total"]
    assert ocr.calls == []


@pytest.mark.parametrize("content_type", [DOCX_CONTENT_TYPE, XLSX_CONTENT_TYPE])
def test_ready_office_delivery_acks_without_work(content_type: str) -> None:
    settings = make_settings()
    document = file_document(status="ready", page_count=2, content_type=content_type, file_name="Готово.docx")
    store = MemoryFiles()
    store.add(document)
    blobs = MemoryBlobs(b"unused")
    pipeline, ocr, embedding, qdrant = build_pipeline(settings, store, blobs)
    result = pipeline.handle(message_for(document, content_type=content_type), redelivered=True)
    assert result.action == "ack"
    assert blobs.reads == []
    assert qdrant.ops == []
    assert ocr.calls == []
    assert embedding.calls == []
    assert embedding.image_calls == []
    assert store.get_by_file_id(document["file_id"])["status"] == "ready"
